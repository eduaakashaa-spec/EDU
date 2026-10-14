"""Assessment results — capture, member view, admin/mentor view, export, import.

Public:
    POST /api/assessments?source=<slug>        ← replaces the Google Apps Script saves

Member (login required):
    GET  /profile/assessments/<id>             ← read their own result

Admin (admin tier — the counselling team):
    GET  /admin/assessments                    ← list / search / filter
    GET  /admin/assessments/<id>               ← full result + link to a member
    POST /admin/assessments/<id>/link          ← attach to a member account (by email)
    POST /admin/assessments/<id>/delete
    GET  /admin/assessments/export.xlsx        ← flat table for analysis
    GET|POST /admin/assessments/import         ← bring in history exported from Google Sheets
"""
import csv
import io
import json
from datetime import datetime, timezone

from flask import (Blueprint, abort, flash, jsonify, redirect, render_template,
                   request, send_file, url_for)
from flask_login import current_user, login_required

from app.decorators import admin_required
from app.extensions import db
from app.models import AssessmentResult, User
from app.services import assessments as svc

assessments_bp = Blueprint('assessments', __name__)


# --------------------------------------------------------------------------- #
# capture
# --------------------------------------------------------------------------- #
@assessments_bp.route('/api/assessments', methods=['POST'])
def submit():
    data = request.get_json(silent=True, force=True)
    if not isinstance(data, dict) or not data:
        return jsonify({'ok': False, 'error': 'Send the assessment as a JSON object.'}), 400
    source = (request.args.get('source') or data.get('source') or '').strip()
    if not source:
        return jsonify({'ok': False, 'error': 'Missing assessment source.'}), 400
    user = current_user if current_user.is_authenticated else None
    row = svc.ingest(source, data, user=user)
    return jsonify({'ok': True, 'id': row.id, 'linked': bool(row.user_id),
                    'view_url': url_for('assessments.member_view', rid=row.id) if row.user_id else None})


# --------------------------------------------------------------------------- #
# member view
# --------------------------------------------------------------------------- #
@assessments_bp.route('/profile/assessments/<int:rid>')
@login_required
def member_view(rid):
    row = db.session.get(AssessmentResult, rid) or abort(404)
    if not svc.can_view(current_user, row):
        abort(404)
    return render_template('members/assessment_detail.html', r=row, s=svc.summarize(row),
                           label=svc.label_for(row.source))


# --------------------------------------------------------------------------- #
# admin
# --------------------------------------------------------------------------- #
def _filtered(args):
    q = (args.get('q') or '').strip()
    source = (args.get('source') or '').strip()
    link = (args.get('link') or '').strip()
    query = AssessmentResult.query
    if q:
        like = f'%{q}%'
        query = query.filter(db.or_(
            AssessmentResult.student_name.ilike(like), AssessmentResult.student_email.ilike(like),
            AssessmentResult.parent_name.ilike(like), AssessmentResult.parent_email.ilike(like),
            AssessmentResult.school.ilike(like), AssessmentResult.result_label.ilike(like)))
    if source:
        query = query.filter(AssessmentResult.source == source)
    if link == 'linked':
        query = query.filter(AssessmentResult.user_id.isnot(None))
    elif link == 'unlinked':
        query = query.filter(AssessmentResult.user_id.is_(None))
    return query.order_by(AssessmentResult.created_at.desc()), q, source, link


@assessments_bp.route('/admin/assessments')
@admin_required
def admin_list():
    query, q, source, link = _filtered(request.args)
    page = request.args.get('page', 1, type=int)
    pagination = query.paginate(page=page, per_page=30, error_out=False)
    totals = dict(db.session.query(AssessmentResult.source, db.func.count(AssessmentResult.id))
                  .group_by(AssessmentResult.source).all())
    total = sum(totals.values())
    linked = AssessmentResult.query.filter(AssessmentResult.user_id.isnot(None)).count()
    return render_template('admin/assessments_list.html', admin_tab='assessments',
                           pagination=pagination, rows=pagination.items, q=q, source=source, link=link,
                           sources=svc.SOURCES, totals=totals, total=total, linked=linked,
                           label_for=svc.label_for)


@assessments_bp.route('/admin/assessments/<int:rid>')
@admin_required
def admin_detail(rid):
    row = db.session.get(AssessmentResult, rid) or abort(404)
    owner = db.session.get(User, row.user_id) if row.user_id else None
    # a registered account with the same email that is not yet linked
    suggestion = None
    if owner is None:
        for em in (row.student_email, row.parent_email):
            if em:
                suggestion = User.query.filter(db.func.lower(User.email) == em.lower()).first()
                if suggestion:
                    break
    return render_template('admin/assessment_detail.html', admin_tab='assessments', r=row,
                           s=svc.summarize(row), label=svc.label_for(row.source),
                           owner=owner, suggestion=suggestion,
                           pretty=json.dumps(row.payload, indent=2, ensure_ascii=False)[:60000])


@assessments_bp.route('/admin/assessments/<int:rid>/link', methods=['POST'])
@admin_required
def admin_link(rid):
    row = db.session.get(AssessmentResult, rid) or abort(404)
    email = (request.form.get('email') or '').strip().lower()
    if not email:
        row.user_id = None
        db.session.commit()
        flash('Result unlinked from any member.', 'success')
    else:
        user = User.query.filter(db.func.lower(User.email) == email).first()
        if user is None:
            flash(f'No member account with the email {email}.', 'danger')
        else:
            row.user_id = user.id
            db.session.commit()
            flash(f'Linked to {user.name} ({user.email}). They can now see it in their profile.', 'success')
    return redirect(url_for('assessments.admin_detail', rid=row.id))


@assessments_bp.route('/admin/assessments/<int:rid>/delete', methods=['POST'])
@admin_required
def admin_delete(rid):
    row = db.session.get(AssessmentResult, rid) or abort(404)
    db.session.delete(row)
    db.session.commit()
    flash('Assessment result deleted.', 'success')
    return redirect(url_for('assessments.admin_list'))


@assessments_bp.route('/admin/assessments/export.xlsx')
@admin_required
def admin_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    query, *_ = _filtered(request.args)
    rows = query.all()

    # one column per score key seen, so the sheet is directly pivotable
    score_keys = []
    for r in rows:
        for k in r.scores:
            if k not in score_keys:
                score_keys.append(k)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Assessments'
    headers = (['ID', 'Submitted (UTC)', 'Assessment', 'Source', 'Member email', 'Member name',
                'Student', 'Student email', 'Student phone', 'Parent', 'Parent email', 'Parent phone',
                'School', 'City', 'Board', 'Class', 'Result', 'Score %', 'Language', 'Origin']
               + [f'Score: {k}' for k in score_keys] + ['Full payload (JSON)'])
    ws.append(headers)
    fill = PatternFill('solid', fgColor='0E3A8A')
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = fill
    members = {u.id: u for u in User.query.filter(User.id.in_({r.user_id for r in rows if r.user_id})).all()} \
        if rows else {}
    for r in rows:
        m = members.get(r.user_id)
        sc = r.scores
        ws.append([r.id, r.created_at.strftime('%Y-%m-%d %H:%M') if r.created_at else '',
                   svc.label_for(r.source), r.source, m.email if m else '', m.name if m else '',
                   r.student_name, r.student_email, r.student_phone, r.parent_name, r.parent_email,
                   r.parent_phone, r.school, r.city, r.board, r.current_class, r.result_label,
                   r.result_score, r.lang, r.origin]
                  + [sc.get(k) for k in score_keys] + [(r.payload_json or '')[:30000]])
    ws.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True,
                     download_name=f'eduaakashaa_assessments_{datetime.now().strftime("%Y%m%d")}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# --------------------------------------------------------------------------- #
# import (history exported from the Google Sheets)
# --------------------------------------------------------------------------- #
_TS_KEYS = ('timestamp', 'submittedat', 'submitted', 'date', 'createdat', 'time')


def _parse_when(row):
    from app.services.assessments import _norm_key
    for k, v in row.items():
        if _norm_key(k) in _TS_KEYS and v:
            s = str(v).strip()
            for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%dT%H:%M:%S', '%d/%m/%Y %H:%M:%S', '%m/%d/%Y %H:%M:%S',
                        '%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
                try:
                    return datetime.strptime(s[:19] if 'T' in fmt or ':' in fmt else s[:10], fmt)
                except ValueError:
                    continue
            try:
                return datetime.fromisoformat(s.replace('Z', '+00:00')).replace(tzinfo=None)
            except ValueError:
                pass
    return None


def _inflate(row):
    """Sheet cells that hold JSON become real nested objects."""
    out = {}
    for k, v in row.items():
        if v is None or str(v).strip() == '':
            continue
        if isinstance(v, str) and v.strip()[:1] in '{[':
            try:
                v = json.loads(v)
            except ValueError:
                pass
        out[str(k).strip()] = v
    return out


@assessments_bp.route('/admin/assessments/import', methods=['GET', 'POST'])
@admin_required
def admin_import():
    result = None
    if request.method == 'POST':
        f = request.files.get('file')
        source = (request.form.get('source') or '').strip()
        if not f or not f.filename:
            flash('Choose a CSV or JSON file exported from the Google Sheet.', 'danger')
            return redirect(url_for('assessments.admin_import'))
        if source not in svc.SOURCES:
            flash('Pick which assessment this sheet belongs to.', 'danger')
            return redirect(url_for('assessments.admin_import'))
        raw = f.read(8_000_000)
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('latin-1')
        try:
            if f.filename.lower().endswith('.json'):
                rows = json.loads(text)
                rows = rows if isinstance(rows, list) else [rows]
            else:
                rows = list(csv.DictReader(io.StringIO(text)))
        except ValueError as exc:
            flash(f'Could not read the file: {exc}', 'danger')
            return redirect(url_for('assessments.admin_import'))

        added = skipped = failed = 0
        for raw_row in rows[:5000]:
            try:
                if not isinstance(raw_row, dict):
                    failed += 1
                    continue
                data = _inflate(raw_row)
                if not data:
                    skipped += 1
                    continue
                ref = svc.content_hash(data)
                before = AssessmentResult.query.filter_by(source=source, source_ref=ref).first()
                if before:
                    skipped += 1
                    continue
                svc.ingest(source, data, origin='import', source_ref=ref,
                           created_at=_parse_when(raw_row), commit=False)
                added += 1
            except Exception:                       # one bad row must not stop the import
                db.session.rollback()
                failed += 1
        db.session.commit()
        result = {'added': added, 'skipped': skipped, 'failed': failed, 'total': len(rows)}
        flash(f'Import finished: {added} added, {skipped} skipped (blank or already imported), {failed} failed.',
              'success' if added else 'info')
    return render_template('admin/assessments_import.html', admin_tab='assessments',
                           sources=svc.SOURCES, result=result)
